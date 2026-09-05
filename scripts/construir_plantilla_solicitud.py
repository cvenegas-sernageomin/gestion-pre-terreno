# -*- coding: utf-8 -*-
"""Construye assets/plantilla_solicitud.xlsx: el formulario oficial "REQUERIMIENTO
CONDUCTOR Y CAMIONETA SERNAGEOMIN" en blanco, con TOKENS @@...@@ sembrados en las
celdas rellenables. Misma estrategia que generador-memos/scripts/construir_plantilla_base.py
(ver memoria: reference-excel-membrete-tokens-jszip-navegador): la plantilla conserva
formato/celdas combinadas exactas; en el navegador se reemplazan los tokens como texto
en el XML del .xlsx (que es un zip), sin tocar xl/media ni xl/drawings.

Uso:
    python construir_plantilla_solicitud.py [origen.xlsx] [salida.xlsx]
"""
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell

RAIZ = Path(__file__).resolve().parents[1]
ORIGEN_DEFAULT = (
    r"C:\Users\carlos.venegas\OneDrive - Sernageomin\05 Central Los Cipreses"
    r"\01 Gestion\02 Preparacion terrenos"
    r"\FORMATO DE REQUERIMIENTO D ECONDUCTOR Y CAMIONETA 24.xlsx"
)
SALIDA_DEFAULT = RAIZ / "assets" / "plantilla_solicitud.xlsx"

MAX_VEHICULOS = 4  # filas disponibles en la plantilla real (7-10, 20-23, 25-28)


def _set(ws, ref, valor):
    if isinstance(ws[ref], MergedCell):
        return False
    ws[ref] = valor
    return True


def sembrar(ws):
    _set(ws, "C4", "@@CENTRO_COSTO@@")
    _set(ws, "E4", "@@REGION_INST@@")
    _set(ws, "G4", "@@FECHA_SOLICITUD@@")
    _set(ws, "C5", "@@JEFE_CC@@")
    _set(ws, "E5", "@@JEFE_PROGRAMA@@")

    for i in range(1, MAX_VEHICULOS + 1):
        fila = 6 + i  # filas 7..10
        _set(ws, f"B{fila}", f"@@PROF_{i}@@")
        _set(ws, f"C{fila}", f"@@COND_{i}@@")
        _set(ws, f"D{fila}", f"@@PAT_{i}@@")
        _set(ws, f"E{fila}", f"@@ORIG_{i}@@")

    _set(ws, "C11", "@@DEPTO@@")
    _set(ws, "C12", "@@OBJETIVO@@")
    _set(ws, "C14", "@@CANT_CAM@@")
    _set(ws, "C15", "@@CARRO_ARRASTRE@@")
    _set(ws, "C16", "@@REGION_LOC@@")
    _set(ws, "E17", "@@ALTURA_SI@@")
    _set(ws, "G17", "@@ALTURA_NO@@")
    _set(ws, "C18", "@@PERMISO_CIRC@@")

    for i in range(1, MAX_VEHICULOS + 1):
        fila = 19 + i  # filas 20..23
        _set(ws, f"B{fila}", f"@@FI_PROF_{i}@@")
        _set(ws, f"D{fila}", f"@@FT_PROF_{i}@@")
        _set(ws, f"F{fila}", f"@@NOM_PROF_{i}@@")

    for i in range(1, MAX_VEHICULOS + 1):
        fila = 24 + i  # filas 25..28
        _set(ws, f"B{fila}", f"@@FI_COND_{i}@@")
        _set(ws, f"D{fila}", f"@@FT_COND_{i}@@")
        _set(ws, f"F{fila}", f"@@NOM_COND_{i}@@")

    _set(ws, "B30", "@@OBSERVACION@@")


def construir(origen, salida):
    wb = openpyxl.load_workbook(origen)
    ws = wb.worksheets[0]
    sembrar(ws)

    salida = Path(salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)

    wb2 = openpyxl.load_workbook(salida)
    ws2 = wb2.worksheets[0]
    print(f"Plantilla con tokens guardada: {salida}")
    print(f"Hoja: {ws2.title!r}  imgs={len(getattr(ws2, '_images', []))}")


if __name__ == "__main__":
    origen = sys.argv[1] if len(sys.argv) > 1 else ORIGEN_DEFAULT
    salida = sys.argv[2] if len(sys.argv) > 2 else SALIDA_DEFAULT
    construir(origen, salida)
