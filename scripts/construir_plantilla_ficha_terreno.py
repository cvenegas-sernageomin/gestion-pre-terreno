# -*- coding: utf-8 -*-
"""Construye assets/plantilla_ficha_terreno.xlsx: la "FICHA PARA COMETIDOS
FUNCIONARIOS" en blanco, con TOKENS @@...@@ sembrados en las celdas rellenables.
Misma estrategia que construir_plantilla_solicitud.py (ver memoria:
reference-excel-membrete-tokens-jszip-navegador).

OJO viáticos: el Excel real calcula E33/F33/G33/E34/F34/G34/C35:G35/E42 con
FÓRMULAS (=O33*C33, etc.) que dependen de las celdas de días. Como el token
queda como texto (t="inlineStr"), esas fórmulas dejarían de funcionar si solo
sembramos las celdas de entrada. Por eso aquí se siembran TAMBIÉN las celdas
de fórmula (se pierden como fórmulas vivas, pero la app JS calcula los mismos
montos con las mismas tarifas oficiales y los escribe ya resueltos).

Uso:
    python construir_plantilla_ficha_terreno.py [origen.xlsx] [salida.xlsx]
"""
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell

RAIZ = Path(__file__).resolve().parents[1]
ORIGEN_DEFAULT = (
    r"C:\Users\carlos.venegas\.claude\uploads\81f12f45-51d7-4e64-9bfd-6b9719b46c17"
    r"\10a2ba5d-Ficha_terreno.xlsx"
)
SALIDA_DEFAULT = RAIZ / "assets" / "plantilla_ficha_terreno.xlsx"

MAX_PARTICIPANTES = 5   # filas 17..21
MAX_ACTIVIDADES = 8     # filas 45..52
MAX_EPP = 10            # filas 55..64


def _set(ws, ref, valor):
    if isinstance(ws[ref], MergedCell):
        return False
    ws[ref] = valor
    return True


def sembrar(ws):
    # --- Identificación del cometido ---
    _set(ws, "C8", "@@SUBDIRECCION@@")
    _set(ws, "G8", "@@DEPTO@@")
    _set(ws, "L8", "@@UNIDAD@@")
    _set(ws, "S8", "@@CENTRO_COSTO@@")
    _set(ws, "C9", "@@OBJETIVO@@")
    _set(ws, "C11", "@@EQUIPO_SI_LABEL@@")
    _set(ws, "D11", "@@EQUIPO_NO_LABEL@@")
    _set(ws, "C12", "@@OBSERVACIONES_GRAL@@")

    # --- Participantes (filas 17..21) ---
    campos_participante = ["NOMBRE", "ORIGEN", "RUT", "CEL", "TEL_EMERG", "CONTACTO_EMERG",
                            "FSALIDA", "FREGRESO", "TRASLADO", "VUELO_IDA", "VUELO_REGRESO",
                            "EX_ALTFIS", "EX_ALTGEO", "EX_FRIO", "EX_CALOR", "EX_CONFIN", "EX_COND"]
    cols_participante = ["B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L",
                          "M", "N", "O", "P", "Q", "S"]
    for i in range(1, MAX_PARTICIPANTES + 1):
        fila = 16 + i  # 17..21
        for campo, col in zip(campos_participante, cols_participante):
            _set(ws, f"{col}{fila}", f"@@P{i}_{campo}@@")

    # --- Destino / localidades ---
    _set(ws, "C23", "@@CIUDAD_DESTINO@@")
    _set(ws, "D24", "@@LOC1@@")
    _set(ws, "H24", "@@LOC2@@")
    _set(ws, "L24", "@@LOC3@@")
    _set(ws, "D25", "@@REF1@@")
    _set(ws, "H25", "@@REF2@@")
    _set(ws, "L25", "@@REF3@@")

    # --- Vehículos ---
    _set(ws, "D27", "@@VEH_CAMIONETA@@")
    _set(ws, "H27", "@@VEH_CARRO@@")
    _set(ws, "L27", "@@VEH_OTRO@@")
    _set(ws, "C28", "@@VALOR_COMBUSTIBLE@@")
    _set(ws, "C29", "@@OBS_TRANSPORTE@@")
    _set(ws, "C30", "@@OTROS_TRANSPORTE@@")

    # --- Viáticos (entrada + celdas que en el original eran fórmulas) ---
    for ref in ["C33", "D33", "E33", "F33", "G33", "C34", "D34", "E34", "F34", "G34",
                "C35", "D35", "E35", "F35", "G35"]:
        _set(ws, ref, f"@@VIA_{ref}@@")

    # --- Fondos a rendir ---
    _set(ws, "C41", "@@FONDOS_OTRO_LABEL@@")
    for ref in ["E38", "E39", "E40", "E41", "E42"]:
        _set(ws, ref, f"@@FONDOS_{ref}@@")
    for ref in ["G38", "G39", "G40", "G41"]:
        _set(ws, ref, f"@@FONDOS_DET_{ref}@@")

    # --- Actividades (filas 45..52) ---
    for i in range(1, MAX_ACTIVIDADES + 1):
        fila = 44 + i  # 45..52
        _set(ws, f"D{fila}", f"@@ACT{i}_DIAS@@")
        _set(ws, f"E{fila}", f"@@ACT{i}_DESC@@")

    # --- EPP (filas 55..64) ---
    for i in range(1, MAX_EPP + 1):
        fila = 54 + i  # 55..64
        _set(ws, f"B{fila}", f"@@EPP{i}_NOMBRE@@")
        _set(ws, f"C{fila}", f"@@EPP{i}_ELEMENTO@@")
        _set(ws, f"D{fila}", f"@@EPP{i}_CANTIDAD@@")
        _set(ws, f"E{fila}", f"@@EPP{i}_ESTADO@@")
        _set(ws, f"F{fila}", f"@@EPP{i}_OBS@@")

    # --- Seguridad y riesgos ---
    _set(ws, "C68", "@@RIESGOS_IDENT@@")
    _set(ws, "C69", "@@RIESGOS_MEDIDAS@@")
    _set(ws, "C70", "@@RIESGOS_PLAN_EMERG@@")

    # --- Fechas y responsables ---
    _set(ws, "C74", "@@ELAB_DIA@@")
    _set(ws, "D74", "@@ELAB_MES@@")
    _set(ws, "E74", "@@ELAB_ANIO@@")
    _set(ws, "C75", "@@ELABORADO_POR@@")
    _set(ws, "C79", "@@REV_DIA@@")
    _set(ws, "D79", "@@REV_MES@@")
    _set(ws, "E79", "@@REV_ANIO@@")
    _set(ws, "C81", "@@REVISADO_POR@@")

    # --- Nota permiso de circulación ---
    _set(ws, "B86", "@@NOTA_PERMISO_CIRCULACION@@")


def construir(origen, salida):
    wb = openpyxl.load_workbook(origen)
    ws = wb.worksheets[0]
    sembrar(ws)

    salida = Path(salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)

    wb2 = openpyxl.load_workbook(salida)
    ws2 = wb2.worksheets[0]
    print(f"Plantilla con tokens guardada: {salida}")
    print(f"Hoja: {ws2.title!r}  imgs={len(getattr(ws2, '_images', []))}")


if __name__ == "__main__":
    origen = sys.argv[1] if len(sys.argv) > 1 else ORIGEN_DEFAULT
    salida = sys.argv[2] if len(sys.argv) > 2 else SALIDA_DEFAULT
    construir(origen, salida)
